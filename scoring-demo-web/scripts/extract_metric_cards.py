"""Build the browser-ready GS/FS rule registry from the two source DOCX files."""

from __future__ import annotations

import argparse
import json
import os
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = ROOT / "scoring-demo-web" / "app" / "data" / "metric-cards.json"


def clean(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def field_map(table) -> dict[str, str]:
    result: dict[str, str] = {}
    for row in table.rows[1:]:
        if len(row.cells) < 2:
            continue
        key = clean(row.cells[0].text)
        value = clean(row.cells[1].text)
        if key:
            result[key] = value
    return result


def dependency_tags(required: str, calculation: str) -> list[str]:
    # Dependencies are taken from the explicit “所需点” contract only.  The
    # calculation prose often mentions trajectories or rackets as context and
    # must not silently turn those into hard model requirements.
    text = required
    tags: list[str] = []
    if re.search(r"J\d{3}|COCO-?17", text, re.I):
        tags.append("pose")
    if re.search(r"BALL", text, re.I):
        tags.append("ball")
    if re.search(r"RK", text, re.I):
        tags.append("racket")
    if re.search(r"场地|球场|单应性", text, re.I):
        tags.append("court")
    if re.search(r"track_id", text, re.I):
        tags.append("tracking")
    return tags


def stage_index(document) -> dict[str, dict[str, str | int]]:
    overview = document.tables[0]
    result: dict[str, dict[str, str | int]] = {}
    for order, row in enumerate(overview.rows[1:], start=1):
        cells = [clean(cell.text) for cell in row.cells]
        if not cells or not cells[0]:
            continue
        result[cells[0]] = {
            "code": cells[0],
            "name": cells[1] if len(cells) > 1 else cells[0],
            "startAction": cells[2] if len(cells) > 2 else "",
            "endAction": cells[3] if len(cells) > 3 else "",
            "order": order,
        }
    return result


def extract_card(domain: str, fields: dict[str, str], stages: dict[str, dict]) -> dict:
    card_id = fields.get("指标编号", "")
    if domain == "GS":
        stage_code = "-".join(card_id.split("-")[:2])
        event_code = card_id.split("-")[0]
    else:
        event_code = card_id.split("-")[0]
        stage_code = event_code
    stage = stages.get(stage_code, {"code": stage_code, "name": stage_code, "startAction": "", "endAction": "", "order": 0})
    required = fields.get("所需点", "")
    calculation = fields.get("计算方式", "")
    return {
        "id": card_id,
        "domain": domain,
        "eventCode": event_code,
        "stageCode": stage_code,
        "stageName": stage["name"],
        "stageOrder": stage["order"],
        "startAction": fields.get("开始标志性动作", "") or stage["startAction"],
        "endAction": fields.get("结束标志性动作", "") or stage["endAction"],
        "name": fields.get("指标名称", ""),
        "definition": fields.get("技术定义", ""),
        "sourceKey": fields.get("原文关键项", ""),
        "reuseSource": fields.get("来源/复用", ""),
        "currentPoints": fields.get("当前识别点", ""),
        "requiredPoints": required,
        "sourceStatus": fields.get("当前状态", ""),
        "calculation": calculation,
        "dependencies": dependency_tags(required, calculation),
        "grades": {
            "A": fields.get("A级", ""),
            "B": fields.get("B级", ""),
            "C": fields.get("C级", ""),
            "D": fields.get("D级", ""),
            "E": fields.get("E级", ""),
        },
        "unavailable": fields.get("不可评价", ""),
        "positiveFeedback": fields.get("AI正向反馈", ""),
        "improvementFeedback": fields.get("AI改进反馈", ""),
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build the browser-ready GS/FS metric-card registry from source DOCX files."
    )
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=os.environ.get("RALLYMATE_METRIC_CARD_SOURCE_DIR"),
        help=(
            "Directory containing the GS and FS source DOCX files. "
            "Alternatively set RALLYMATE_METRIC_CARD_SOURCE_DIR."
        ),
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if args.source_dir is None:
        parser.error(
            "--source-dir is required when RALLYMATE_METRIC_CARD_SOURCE_DIR is unset"
        )

    source_dir = args.source_dir.expanduser().resolve()
    output = args.output.expanduser().resolve()
    source_files = {
        "GS": next(source_dir.glob("GS*.docx")),
        "FS": next(source_dir.glob("FS*.docx")),
    }
    cards: list[dict] = []
    sources: list[dict] = []
    for domain, source in source_files.items():
        document = Document(source)
        stages = stage_index(document)
        domain_cards = [
            extract_card(domain, field_map(table), stages)
            for table in document.tables[1:]
        ]
        domain_cards = [card for card in domain_cards if card["id"].startswith(domain)]
        cards.extend(domain_cards)
        sources.append(
            {
                "domain": domain,
                "fileName": source.name,
                "indicatorCount": len(domain_cards),
                "stageCount": len(stages),
            }
        )

    if len(cards) != 298:
        raise RuntimeError(f"Expected 298 cards, found {len(cards)}")

    payload = {
        "schemaVersion": "1.0.0",
        "registryVersion": "2026-08-10-demo.1",
        "generatedAt": "2026-08-10",
        "sources": sources,
        "cards": cards,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {len(cards)} cards to {output}")
    for domain in ("GS", "FS"):
        subset = [card for card in cards if card["domain"] == domain]
        dep_counts = {
            dep: sum(dep in card["dependencies"] for card in subset)
            for dep in ("pose", "ball", "racket", "court", "tracking")
        }
        print(domain, len(subset), dep_counts)


if __name__ == "__main__":
    main()
